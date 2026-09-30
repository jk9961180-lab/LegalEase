
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


def create_txt(text: str) -> bytes:
    return text.encode("utf-8")


def create_docx(text: str, document_type: str) -> bytes:
    document = Document()
    section = document.sections[0]

    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(document_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for paragraph_text in text.split("\n\n"):
        paragraph_text = paragraph_text.strip()
        if not paragraph_text:
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.15
        run = paragraph.add_run(paragraph_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(
        "Generated with LegalEase - AI-Powered Legal Document Generator"
    ).font.size = Pt(8)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


class LegalEasePDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 10, "LegalEase", align="C")
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "", 8)
        self.cell(
            0, 10,
            "Generated with LegalEase - AI-generated draft",
            align="C"
        )


def create_pdf(text: str, document_type: str) -> bytes:
    pdf = LegalEasePDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, document_type.upper(), align="C")
    pdf.ln(5)

    pdf.set_font("Helvetica", "", 11)

    for paragraph in text.split("\n\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            continue

        safe_paragraph = (
            paragraph.encode("latin-1", "replace").decode("latin-1")
        )
        pdf.multi_cell(0, 7, safe_paragraph)
        pdf.ln(3)

    return bytes(pdf.output())